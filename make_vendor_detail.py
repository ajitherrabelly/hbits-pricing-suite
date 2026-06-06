# -*- coding: utf-8 -*-
"""Extract contract-level vendor closure detail into vendor_detail.js (separate from analytics)."""
import openpyxl, os, json, datetime
from make_analytics import DETAIL, find_header, canon

OUT = "vendor_detail.js"


def fmt_date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%m/%d/%Y")
    return str(v).strip() if v else ""


def num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def read_full(path, year):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    target = None
    for sn in wb.sheetnames:
        ws = wb[sn]
        hi, hdr = find_header(ws)
        if hi is not None:
            target = (sn, hi, hdr); break
    sn, hi, hdr = target
    ws = wb[sn]
    col = {}
    for j, c in enumerate(hdr):
        if "VENDOR NAME" in c: col["v"] = j
        elif "DEPARTMENT" in c or "FACILITY" in c: col["dept"] = j
        elif "CONTRACT NUMBER" in c: col["no"] = j
        elif "CURRENT CONTRACT" in c or "CONTRACT AMOUN" in c: col["amt"] = j
        elif "SPENDING" in c: col["spent"] = j
        elif "START" in c: col["start"] = j
        elif "END" in c: col["end"] = j
        elif "DESCRIPTION" in c: col["desc"] = j
    def get(row, key):
        j = col.get(key)
        return row[j] if j is not None and j < len(row) else None
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= hi:
            continue
        v = get(row, "v")
        if not v or str(v).strip().lower() in ("grand total", "total", "row labels"):
            continue
        out.append({
            "v": canon(v), "y": year,
            "dept": (str(get(row, "dept")).strip() if get(row, "dept") else "Unknown"),
            "no": (str(get(row, "no")).strip() if get(row, "no") else ""),
            "amt": round(num(get(row, "amt"))),
            "spent": round(num(get(row, "spent"))),
            "start": fmt_date(get(row, "start")),
            "end": fmt_date(get(row, "end")),
            "desc": (str(get(row, "desc")).strip()[:70] if get(row, "desc") else ""),
        })
    wb.close()
    return out


def main():
    years = [2023, 2024, 2025]
    contracts = []
    for y in years:
        contracts += read_full(DETAIL[y][0], y)

    vy, vd = {}, {}
    for c in contracts:
        vy.setdefault(c["v"], {}).setdefault(c["y"], 0)
        vy[c["v"]][c["y"]] += 1
        vd.setdefault(c["v"], {}).setdefault(c["y"], 0)
        vd[c["v"]][c["y"]] += c["amt"]

    vendors = []
    for v, yc in vy.items():
        vendors.append({
            "name": v,
            "c2023": yc.get(2023, 0), "c2024": yc.get(2024, 0), "c2025": yc.get(2025, 0),
            "total": sum(yc.values()),
            "d2023": vd[v].get(2023, 0), "d2024": vd[v].get(2024, 0), "d2025": vd[v].get(2025, 0),
            "dtotal": sum(vd[v].values()),
        })
    vendors.sort(key=lambda r: r["total"], reverse=True)

    data = {
        "years": years,
        "vendors": vendors,
        "contracts": contracts,
        "totals": {
            "closures": len(contracts),
            "dollars": sum(c["amt"] for c in contracts),
            "vendors": len(vendors),
        },
        "generated": "from HBITS per-year closure master workbooks (Jan 2023 - Dec 2025)",
    }
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("window.VENDOR_DETAIL = " + json.dumps(data) + ";\n")

    print(f"contracts: {len(contracts)} | vendors: {len(vendors)} | ${data['totals']['dollars']:,}")
    print(f"per year: " + ", ".join(f"{y}={sum(1 for c in contracts if c['y']==y)}" for y in years))
    print("top5:", [(v["name"], v["total"]) for v in vendors[:5]])
    print("wrote", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
