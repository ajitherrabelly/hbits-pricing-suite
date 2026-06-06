# -*- coding: utf-8 -*-
"""Extract executive analytics from the HBITS closure/award workbooks into analytics_data.js."""
import openpyxl, os, re, json

BASE = "drive-download-20260606T205412Z-3-001"
OUT = "analytics_data.js"

DETAIL = {
    2023: (f"{BASE}/Vendor Closure Data of year 2023/Final Filter data 2023.xlsx", None),
    2024: (f"{BASE}/Vendor Closure Data of year 2024/All companies No_s of Closures 2024.xlsx", None),
    2025: (f"{BASE}/Vendor Closure Data of year 2025/All companies Closures 2025.xlsx", None),
}
TOP10 = {
    2023: f"{BASE}/Top 10 Titles & Vendors of Year 2023.xlsx",
    2024: f"{BASE}/Top 10 titles & Vendors of year 2024.xlsx",
    2025: f"{BASE}/Top 10 Titles_s & Vendor_s of year 2025.xlsx",
}

# canonical display names keyed by a normalized token
NICE = {
    "AVENUES": "Avenues International", "BITSBYTES": "Bits & Bytes", "BNB": "Bits & Bytes",
    "BROADCROSS": "Broad Crossing", "COMPUTERTE": "Computer Technology Svc (CTS)", "CTS": "Computer Technology Svc (CTS)",
    "CROSSFIRE": "Crossfire Consulting", "CURRIER": "CMA Consulting (Currier McCabe)", "CMA": "CMA Consulting (Currier McCabe)",
    "EXPERIS": "Experis US", "GCOM": "GCOM Software / Voyatek", "GENESYS": "Genesys Consulting", "GENSYS": "Genesys Consulting",
    "GREYCELL": "Greycell Labs", "ILINK": "I-Link Solutions", "JSM": "JSM Consulting",
    "KNOWLEDGE": "Knowledge Builders (KBI)", "KBI": "Knowledge Builders (KBI)", "LANCESOFT": "LanceSoft",
    "MINDLANCE": "Mindlance", "MONTCO": "Montco / Rotator", "MONTACO": "Montco / Rotator", "MVP": "MVP Consulting Plus",
    "OST": "OST Inc", "PANHA": "Panha Solutions", "PSI": "PSI International", "RMS": "RMS Computer",
    "SEVEN": "Seven Seas / S2Tech", "SLIGO": "Sligo Software", "SOFTWAREPE": "Software People",
    "SPRUCE": "Spruce Technology", "SVAM": "SVAM International", "SYSTEMEDGE": "System Edge USA",
    "TECHVALLEY": "Tech Valley Talent", "TEKSYSTEMS": "TEKsystems", "TEK": "TEKsystems",
    "TRIGYN": "Trigyn Technologies", "UNIQUE": "Unique Comp", "VGROUP": "V Group",
}


def norm(name):
    if not name:
        return None
    s = re.sub(r"[^A-Za-z0-9]", "", str(name).upper())
    return s


def canon(name):
    s = norm(name)
    if not s:
        return None
    for key in sorted(NICE, key=len, reverse=True):
        if s.startswith(key):
            return NICE[key]
    # fallback: title-case the cleaned words
    return re.sub(r"\s+", " ", re.sub(r"[^A-Za-z0-9 ]", " ", str(name))).strip().title()


def find_header(ws, must="VENDOR NAME"):
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        cells = [str(c).strip().upper() if c is not None else "" for c in row]
        if any(must in c for c in cells):
            return i, cells
    return None, None


def read_detail(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # choose the sheet that has the detail header
    target = None
    for sn in wb.sheetnames:
        ws = wb[sn]
        hi, hdr = find_header(ws)
        if hi is not None:
            target = (sn, hi, hdr); break
    sn, hi, hdr = target
    ws = wb[sn]
    col = {}
    for j, c in enumerate(hdr):
        if "VENDOR NAME" in c: col["vendor"] = j
        elif "DEPARTMENT" in c or "FACILITY" in c: col["agency"] = j
        elif "CURRENT CONTRACT" in c or c == "AMOUNT" or "CONTRACT AMOUN" in c: col["amount"] = j
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= hi:
            continue
        v = row[col["vendor"]] if col.get("vendor") is not None and col["vendor"] < len(row) else None
        if not v or str(v).strip().lower() in ("grand total", "total", "row labels"):
            continue
        ag = row[col["agency"]] if col.get("agency") is not None and col["agency"] < len(row) else None
        amt = 0.0
        if col.get("amount") is not None and col["amount"] < len(row):
            try: amt = float(row[col["amount"]] or 0)
            except (TypeError, ValueError): amt = 0.0
        rows.append((canon(v), (str(ag).strip() if ag else "Unknown"), amt))
    wb.close()
    return rows


def read_titles(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = None
    for sn in wb.sheetnames:
        if sn.lower().startswith("list of all titles"):
            sheet = sn; break
    out = []
    if sheet:
        ws = wb[sheet]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue
            t = row[0] if row else None
            n = row[1] if len(row) > 1 else None
            if t and n is not None and str(t).strip().lower() != "total":
                try: out.append((str(t).strip(), int(float(n))))
                except (TypeError, ValueError): pass
    wb.close()
    return out


def main():
    years = [2023, 2024, 2025]
    vendor_year = {}   # canon -> {year: count}
    vendor_dollars = {}  # canon -> {year: dollars}
    agency_tot = {}    # agency -> {closures, dollars}
    kpis = {}
    for y in years:
        rows = read_detail(DETAIL[y][0])
        kpis[y] = {"closures": len(rows), "dollars": 0.0, "vendors": set()}
        for vendor, agency, amt in rows:
            vendor_year.setdefault(vendor, {}).setdefault(y, 0)
            vendor_year[vendor][y] += 1
            vendor_dollars.setdefault(vendor, {}).setdefault(y, 0.0)
            vendor_dollars[vendor][y] += amt
            kpis[y]["dollars"] += amt
            kpis[y]["vendors"].add(vendor)
            a = agency_tot.setdefault(agency, {"closures": 0, "dollars": 0.0})
            a["closures"] += 1; a["dollars"] += amt
        kpis[y]["vendors"] = len(kpis[y]["vendors"])
        kpis[y]["avg"] = kpis[y]["dollars"] / kpis[y]["closures"] if kpis[y]["closures"] else 0

    vendors = []
    for v, yc in vendor_year.items():
        total = sum(yc.values())
        vendors.append({
            "name": v,
            "c2023": yc.get(2023, 0), "c2024": yc.get(2024, 0), "c2025": yc.get(2025, 0),
            "total": total,
            "d2023": round(vendor_dollars.get(v, {}).get(2023, 0)),
            "d2024": round(vendor_dollars.get(v, {}).get(2024, 0)),
            "d2025": round(vendor_dollars.get(v, {}).get(2025, 0)),
            "dtotal": round(sum(vendor_dollars.get(v, {}).values())),
        })
    vendors.sort(key=lambda r: r["total"], reverse=True)

    titles = {y: read_titles(TOP10[y]) for y in years}

    agencies = [{"name": k, "closures": v["closures"], "dollars": round(v["dollars"])}
                for k, v in agency_tot.items()]
    agencies.sort(key=lambda r: r["closures"], reverse=True)

    total_closures = sum(kpis[y]["closures"] for y in years)
    total_dollars = sum(kpis[y]["dollars"] for y in years)

    data = {
        "years": years,
        "kpi": {
            "perYear": {str(y): {"closures": kpis[y]["closures"], "dollars": round(kpis[y]["dollars"]),
                                  "vendors": kpis[y]["vendors"], "avg": round(kpis[y]["avg"])} for y in years},
            "totalClosures": total_closures,
            "totalDollars": round(total_dollars),
            "activeVendors": len(vendors),
            "avgContract": round(total_dollars / total_closures) if total_closures else 0,
        },
        "vendors": vendors,
        "titles": {str(y): titles[y] for y in years},
        "agencies": agencies[:15],
        "generated": "from HBITS closure/award workbooks (Jan 2023 - Dec 2025)",
    }

    with open(OUT, "w", encoding="utf-8") as f:
        f.write("window.ANALYTICS = " + json.dumps(data, indent=1) + ";\n")

    # ---- console validation ----
    print("Year totals (their pivots: 2023=560, 2024=506):")
    for y in years:
        k = kpis[y]
        print(f"  {y}: {k['closures']} closures | ${round(k['dollars']):,} | {k['vendors']} vendors | avg ${round(k['avg']):,}")
    print(f"3-yr: {total_closures} closures | ${round(total_dollars):,} | {len(vendors)} vendors")
    print("Top 8 vendors by total closures:")
    for v in vendors[:8]:
        print(f"  {v['name']:32s} {v['c2023']:>4}/{v['c2024']:>4}/{v['c2025']:>4} = {v['total']:>4}  ${v['dtotal']:,}")
    print("Top 5 agencies:", [(a['name'][:30], a['closures']) for a in agencies[:5]])
    print("Title demand counts per year:", {y: len(titles[y]) for y in years})
    print("wrote", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
